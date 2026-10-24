"""Shared helpers for all step reports (plain black-and-white Word output)."""
import io
import json
import os
import re
import sys
from datetime import datetime

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 9, "axes.grid": True,
                     "grid.color": "0.85", "axes.spines.top": False, "axes.spines.right": False})

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def new_doc(title, subtitle, simulated=False):
    d = Document()
    st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
    for n in ("Heading 1", "Heading 2", "Title"):
        h = d.styles[n]; h.font.name = "Calibri"; h.font.color.rgb = RGBColor(0, 0, 0)
    for s in d.sections:
        s.left_margin = s.right_margin = Inches(1); s.top_margin = s.bottom_margin = Inches(1)
    d.add_heading(title, 0)
    d.add_paragraph(f"{subtitle}\nGenerated {datetime.now():%Y-%m-%d %H:%M}")
    if simulated:
        p = d.add_paragraph(); r = p.add_run("WARNING: GENERATED FROM SIMULATED / TEST DATA (pipeline test). "
                                             "Do not use any number in the thesis or paper."); r.bold = True
    return d


def table(d, header, rows, widths=None, size=9.5):
    t = d.add_table(rows=1, cols=len(header)); t.style = "Table Grid"
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(str(h)); r.bold = True; r.font.size = Pt(size)
    for row in rows:
        cells = t.add_row().cells
        for i, x in enumerate(row):
            cells[i].text = ""; cells[i].paragraphs[0].add_run(str(x)).font.size = Pt(size)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    return t


def caption(d, text):
    p = d.add_paragraph(); r = p.add_run(text); r.italic = True; r.font.size = Pt(9.5)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def picture(d, path, cap, width=6.0):
    d.add_picture(path, width=Inches(width)); caption(d, cap)


def savefig(path):
    plt.tight_layout(); plt.savefig(path, dpi=200); plt.close()


def dump_json(obj, path):
    with open(path, "w") as f:
        json.dump(obj, f, indent=2, default=float)
