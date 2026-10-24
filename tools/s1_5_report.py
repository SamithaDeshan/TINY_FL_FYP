"""Step S1.5 report generator: CSV session -> thesis-ready DOCX + figures + summary JSON.

    python tools/s1_5_report.py data/s1_5_session1.csv --out reports/s1_5
Labels the report understands (set with 'L <name>' while logging):
    warmup, dark, ambient, bright, still, motion, ramp   (any others are listed but not checked)
"""
import argparse
import io
import json
import os
import re
from datetime import datetime

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 9, "axes.grid": True,
                     "grid.color": "0.85", "axes.spines.top": False, "axes.spines.right": False})

# Engineering acceptance criteria (stated in the report so they can be defended/changed)
CRIT = {"ldr_sd_max_pct_fs": 2.0, "ldr_sep_sd": 10.0, "dht_success_min": 95.0,
        "pir_still_max": 5.0, "pir_motion_min": 10.0, "period_ms": (90, 110),
        "sat_max_pct": 1.0}


def load(path):
    meta, rows = [], []
    with open(path, encoding="utf-8", errors="replace") as f:
        for ln in f:
            ln = ln.rstrip("\r\n")
            if ln.startswith("#"):
                meta.append(ln[1:].strip())
            elif ln and not ln.startswith("ms,"):
                rows.append(ln)
    cols = ["ms", "label", "pir", "ldr_raw", "ldr10", "temp_c", "hum_pct", "dht_status", "pwm"]
    df = pd.read_csv(io.StringIO("\n".join(rows)), names=cols, na_values=["nan"], on_bad_lines="skip")
    df = df.dropna(subset=["ms", "pir", "ldr_raw"]).copy()
    for c in ["ms", "pir", "ldr_raw", "ldr10", "dht_status", "pwm"]:
        df[c] = pd.to_numeric(df[c], errors="coerce")
    df = df.dropna(subset=["ms", "pir", "ldr_raw", "ldr10", "dht_status", "pwm"])
    df["t"] = (df["ms"] - df["ms"].iloc[0]) / 1000.0
    kv = {}
    for m in meta:
        for k, v in re.findall(r"(\w+)=(\S+)", m):
            kv[k] = v
    return df, meta, kv


def fig_to(path):
    plt.tight_layout()
    plt.savefig(path, dpi=200)
    plt.close()


def mark_labels(ax, df):
    ch = df.index[df["label"] != df["label"].shift()]
    ymax = ax.get_ylim()[1]
    for i in ch:
        ax.axvline(df.loc[i, "t"], color="0.5", lw=0.7, ls="--")
        ax.text(df.loc[i, "t"], ymax, df.loc[i, "label"], rotation=90, va="top", ha="right", fontsize=7)


def analyse(df, kv, figdir):
    res, checks, figs = {}, [], {}
    labels = list(dict.fromkeys(df["label"]))
    res["n_rows"] = int(len(df))
    res["duration_s"] = float(df["t"].iloc[-1])
    dt = np.diff(df["ms"].values)
    res["period_ms_median"] = float(np.median(dt))
    res["period_ms_p99"] = float(np.percentile(dt, 99))

    # ---- LDR ----
    g = df.groupby("label")
    ldr = g["ldr_raw"].agg(["count", "mean", "std", "min", "max"]).reindex(labels)
    ldr10 = g["ldr10"].agg(["mean"]).reindex(labels)
    res["ldr_by_label"] = {l: {"n": int(ldr.loc[l, "count"]), "mean_raw": float(ldr.loc[l, "mean"]),
                               "sd_raw": float(ldr.loc[l, "std"]), "min_raw": int(ldr.loc[l, "min"]),
                               "max_raw": int(ldr.loc[l, "max"]),
                               "mean_10bit": float(ldr10.loc[l, "mean"])} for l in labels}
    fig, ax = plt.subplots(figsize=(6.5, 2.6))
    ax.plot(df["t"], df["ldr_raw"], color="black", lw=0.6)
    ax.set_xlabel("Time (s)"); ax.set_ylabel("LDR ADC (12-bit)")
    mark_labels(ax, df)
    p = os.path.join(figdir, "fig_ldr_time.png"); fig_to(p); figs["ldr"] = p
    L = res["ldr_by_label"]
    if all(k in L for k in ("dark", "ambient", "bright")):
        d, a, b = (L[k]["mean_raw"] for k in ("dark", "ambient", "bright"))
        sd = max(L[k]["sd_raw"] for k in ("dark", "ambient", "bright")) or 1.0
        mono_up = d < a < b
        mono_dn = d > a > b
        res["ldr_polarity"] = ("higher ADC = brighter" if mono_up else
                               "higher ADC = darker (inverted divider)" if mono_dn else "not monotonic")
        sep = min(abs(a - d), abs(b - a)) / sd
        res["ldr_separation_sd"] = float(sep)
        checks.append(("LDR separates dark / ambient / bright",
                       (mono_up or mono_dn) and sep >= CRIT["ldr_sep_sd"],
                       f"{res['ldr_polarity']}; min step between levels = {sep:.1f} sd"))
        worst = max(L[k]["sd_raw"] for k in ("dark", "ambient", "bright")) / 4095 * 100
        checks.append(("LDR noise within a steady label",
                       worst <= CRIT["ldr_sd_max_pct_fs"],
                       f"worst SD = {worst:.2f}% of full scale (limit {CRIT['ldr_sd_max_pct_fs']}%)"))
        sat = df[df["label"].isin(["dark", "ambient", "bright"])]
        satp = float(((sat["ldr_raw"] >= 4090) | (sat["ldr_raw"] <= 5)).mean() * 100)
        res["ldr_saturation_pct"] = satp
        checks.append(("LDR not saturated", satp <= CRIT["sat_max_pct"],
                       f"{satp:.1f}% of samples at an ADC rail (limit {CRIT['sat_max_pct']}%)"))
    else:
        checks.append(("LDR dark/ambient/bright test", None, "labels dark, ambient, bright not all present"))

    # ---- PIR ----
    pir = g["pir"].mean().reindex(labels) * 100
    edges = {}
    for l in labels:
        s = df.loc[df["label"] == l, "pir"].values
        edges[l] = int(((s[1:] == 1) & (s[:-1] == 0)).sum()) if len(s) > 1 else 0
    res["pir_by_label"] = {l: {"high_pct": float(pir[l]), "rising_edges": edges[l]} for l in labels}
    first_high = df.loc[df["pir"] == 1, "t"]
    res["pir_first_high_s"] = float(first_high.iloc[0]) if len(first_high) else None
    fig, ax = plt.subplots(figsize=(6.5, 1.8))
    ax.step(df["t"], df["pir"], color="black", lw=0.8, where="post")
    ax.set_ylim(-0.1, 1.3); ax.set_yticks([0, 1]); ax.set_xlabel("Time (s)"); ax.set_ylabel("PIR")
    mark_labels(ax, df)
    p = os.path.join(figdir, "fig_pir_time.png"); fig_to(p); figs["pir"] = p
    P = res["pir_by_label"]
    if "still" in P and "motion" in P:
        checks.append(("PIR quiet when still", P["still"]["high_pct"] <= CRIT["pir_still_max"],
                       f"HIGH {P['still']['high_pct']:.1f}% of time (limit {CRIT['pir_still_max']}%)"))
        checks.append(("PIR responds to motion",
                       P["motion"]["high_pct"] >= CRIT["pir_motion_min"] and P["motion"]["rising_edges"] >= 1,
                       f"HIGH {P['motion']['high_pct']:.1f}%, {P['motion']['rising_edges']} rising edge(s)"))
    else:
        checks.append(("PIR still/motion test", None, "labels still and motion not both present"))

    # ---- DHT11 ----
    att = df[df["dht_status"] != 0]
    ok = int((att["dht_status"] == 1).sum())
    n_att = int(len(att))
    rate = ok / n_att * 100 if n_att else float("nan")
    v = df[df["dht_status"] == 1][["temp_c", "hum_pct"]]
    res["dht"] = {"attempts": n_att, "ok": ok, "success_pct": rate,
                  "temp_mean": float(v["temp_c"].mean()) if len(v) else None,
                  "temp_sd": float(v["temp_c"].std()) if len(v) else None,
                  "temp_min": float(v["temp_c"].min()) if len(v) else None,
                  "temp_max": float(v["temp_c"].max()) if len(v) else None,
                  "hum_mean": float(v["hum_pct"].mean()) if len(v) else None,
                  "hum_sd": float(v["hum_pct"].std()) if len(v) else None,
                  "hum_min": float(v["hum_pct"].min()) if len(v) else None,
                  "hum_max": float(v["hum_pct"].max()) if len(v) else None}
    fig, ax = plt.subplots(figsize=(6.5, 2.4))
    ax.plot(df.loc[df["dht_status"] == 1, "t"], v["temp_c"], "k-o", ms=2, lw=0.7, label="Temperature (C)")
    ax2 = ax.twinx(); ax2.grid(False)
    ax2.plot(df.loc[df["dht_status"] == 1, "t"], v["hum_pct"], color="0.45", ls="--", marker="s", ms=2, lw=0.7,
             label="Humidity (%)")
    ax.set_xlabel("Time (s)"); ax.set_ylabel("Temperature (C)"); ax2.set_ylabel("Humidity (%)")
    h1, l1 = ax.get_legend_handles_labels(); h2, l2 = ax2.get_legend_handles_labels()
    ax.legend(h1 + h2, l1 + l2, loc="best", fontsize=7, frameon=False)
    p = os.path.join(figdir, "fig_dht_time.png"); fig_to(p); figs["dht"] = p
    checks.append(("DHT11 read success", rate >= CRIT["dht_success_min"] if n_att else None,
                   f"{ok}/{n_att} valid ({rate:.1f}%; limit {CRIT['dht_success_min']}%)"))
    if len(v):
        plaus = bool(v["temp_c"].between(0, 50).all() and v["hum_pct"].between(20, 90).all())
        checks.append(("DHT11 values within datasheet range", plaus,
                       f"T {v['temp_c'].min():.0f}-{v['temp_c'].max():.0f} C, "
                       f"RH {v['hum_pct'].min():.0f}-{v['hum_pct'].max():.0f} %"))

    # ---- PWM ramp ----
    r = df[df["label"] == "ramp"]
    if len(r) > 20:
        corr = float(np.corrcoef(r["pwm"], r["ldr_raw"])[0, 1])
        res["ramp_corr_pwm_ldr"] = corr
        fig, ax = plt.subplots(figsize=(3.4, 2.6))
        ax.scatter(r["pwm"], r["ldr_raw"], s=4, color="black")
        ax.set_xlabel("LED PWM duty (0-255)"); ax.set_ylabel("LDR ADC (12-bit)")
        p = os.path.join(figdir, "fig_ramp.png"); fig_to(p); figs["ramp"] = p
        checks.append(("PWM ramp output present", bool(r["pwm"].max() >= 250),
                       f"max duty logged = {int(r['pwm'].max())}; PWM-vs-LDR correlation r = {corr:+.2f} "
                       f"(only meaningful if the LED shines on the LDR)"))

    # ---- timing + OLED ----
    lo, hi = CRIT["period_ms"]
    checks.append(("Logging period 100 ms", lo <= res["period_ms_median"] <= hi,
                   f"median {res['period_ms_median']:.0f} ms, 99th pct {res['period_ms_p99']:.0f} ms"))
    checks.append(("OLED detected at 0x3C", kv.get("oled") == "ok", f"firmware reported oled={kv.get('oled','?')}"))
    return res, checks, figs, labels


# ---------------- DOCX helpers ----------------
def set_style(doc):
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
    for n in ("Heading 1", "Heading 2", "Title"):
        h = doc.styles[n]; h.font.name = "Calibri"; h.font.color.rgb = RGBColor(0, 0, 0)
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Inches(1); sec.top_margin = sec.bottom_margin = Inches(1)


def table(doc, header, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(header)); t.style = "Table Grid"
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(9.5)
    for row in rows:
        cells = t.add_row().cells
        for i, x in enumerate(row):
            cells[i].text = ""; cells[i].paragraphs[0].add_run(str(x)).font.size = Pt(9.5)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    return t


def caption(doc, text):
    p = doc.add_paragraph(); r = p.add_run(text); r.italic = True; r.font.size = Pt(9.5)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def build_doc(path, src, df, meta, kv, res, checks, figs, labels):
    d = Document(); set_style(d)
    d.add_heading("Step S1.5 Report: Hardware Bring-up and Sensor Characterisation", 0)
    d.add_paragraph(f"TinyFL / Stage 1 hardware validation, ESP32-S3 prototype\n"
                    f"Generated {datetime.now():%Y-%m-%d %H:%M} from {os.path.basename(src)}")
    if kv.get("source", "").upper() == "SIMULATED":
        p = d.add_paragraph(); r = p.add_run("WARNING: THIS REPORT WAS GENERATED FROM SIMULATED DATA (pipeline test). "
                                             "Do not use any number in the thesis or paper."); r.bold = True
    d.add_heading("1. Purpose", 1)
    d.add_paragraph("This step verifies that every peripheral on the prototype works at the pin assignments used in the "
                    "firmware, and characterises the raw behaviour of each sensor (noise, range, polarity, response) "
                    "before any machine-learning model is run on real readings. The measurements form the sensor "
                    "characterisation section of the implementation chapter and give the real-sensor input ranges "
                    "needed to validate the simulated training data used in Stage 1 and Stage 2.")
    d.add_heading("2. Test setup", 1)
    table(d, ["Function", "Device", "GPIO"], [
        ["I2C SDA / SCL", "SSD1306 OLED 128x64 (0x3C)", f"{kv.get('oled_sda','8')} / {kv.get('oled_scl','9')}"],
        ["Motion", "PIR (HC-SR501)", kv.get("pir", "4")],
        ["Ambient light", "LDR (GL5528) divider, ADC1", kv.get("ldr", "1")],
        ["Temperature / humidity", "DHT11", kv.get("dht", "5")],
        ["Light actuator", "LED via PWM (5 kHz, 8-bit)", kv.get("led_pwm", "6")]], [1.8, 3.0, 1.0])
    caption(d, "Table 1. Pin assignment used in this session.")
    d.add_paragraph(f"Board: {kv.get('chip','?')} rev {kv.get('rev','?')}, {kv.get('cores','?')} cores at "
                    f"{kv.get('cpu_mhz','?')} MHz. Firmware image: {kv.get('sketch_bytes','?')} bytes (bring-up "
                    f"firmware including OLED and logging libraries, not the ML model). "
                    f"Free heap at boot: {kv.get('heap_free','?')} bytes.")
    d.add_heading("3. Method", 1)
    d.add_paragraph("A single continuous session was logged at 10 Hz (DHT11 sampled every 2 s, its maximum rate being "
                    "1 Hz). The operator set a label before each condition: dark (LDR covered), ambient (room light), "
                    "bright (torch at the LDR), still (no movement in front of the PIR), motion (repeated movement), "
                    "and ramp (LED duty 0 to 255 to 0 over 10 s). "
                    f"The session contains {res['n_rows']} samples over {res['duration_s']:.0f} s. "
                    "The ADC runs at 12 bits (0-4095); a 10-bit value (raw >> 2) is also logged because the Stage 1 "
                    "model was trained on a 0-1023 scale.")
    d.add_heading("4. Results", 1)
    d.add_heading("4.1 Light sensor (LDR)", 2)
    table(d, ["Label", "n", "Mean (12-bit)", "SD", "Min", "Max", "Mean (10-bit)"],
          [[l, v["n"], f"{v['mean_raw']:.0f}", f"{v['sd_raw']:.1f}", v["min_raw"], v["max_raw"],
            f"{v['mean_10bit']:.0f}"] for l, v in res["ldr_by_label"].items()])
    caption(d, "Table 2. LDR readings per labelled condition.")
    d.add_picture(figs["ldr"], width=Inches(6.2)); caption(d, "Figure 1. LDR ADC value over the session.")
    if "ldr_polarity" in res:
        d.add_paragraph(f"Polarity: {res['ldr_polarity']}. The Stage 1 rules (LDR > 700 means bright, LDR < 300 means "
                        "dark, on a 0-1023 scale) hold on this hardware only if the 10-bit means above fall on the "
                        "correct sides of those thresholds; compare Table 2 against them before reusing the model.")
    d.add_heading("4.2 Motion sensor (PIR)", 2)
    table(d, ["Label", "HIGH (% of samples)", "Rising edges"],
          [[l, f"{v['high_pct']:.1f}", v["rising_edges"]] for l, v in res["pir_by_label"].items()])
    caption(d, "Table 3. PIR output per labelled condition.")
    d.add_picture(figs["pir"], width=Inches(6.2)); caption(d, "Figure 2. PIR output over the session.")
    d.add_heading("4.3 Temperature and humidity (DHT11)", 2)
    h = res["dht"]
    f = lambda x, n=1: "n/a" if x is None else f"{x:.{n}f}"
    table(d, ["Quantity", "Mean", "SD", "Min", "Max"],
          [["Temperature (C)", f(h["temp_mean"]), f(h["temp_sd"], 2), f(h["temp_min"], 0), f(h["temp_max"], 0)],
           ["Humidity (%)", f(h["hum_mean"]), f(h["hum_sd"], 2), f(h["hum_min"], 0), f(h["hum_max"], 0)]])
    caption(d, f"Table 4. DHT11 statistics from {h['ok']} valid readings ({h['attempts']} attempts, "
               f"{h['success_pct']:.1f}% success).")
    d.add_picture(figs["dht"], width=Inches(6.2)); caption(d, "Figure 3. DHT11 temperature and humidity.")
    d.add_paragraph("The DHT11 reports integer values with a datasheet accuracy of about +/-2 C and +/-5 %RH, which is "
                    "coarser than the DHT22 named in the Stage 1 paper.")
    if "ramp" in figs:
        d.add_heading("4.4 PWM output", 2)
        d.add_picture(figs["ramp"], width=Inches(3.2))
        caption(d, f"Figure 4. LDR response during the PWM ramp (r = {res['ramp_corr_pwm_ldr']:+.2f}).")
    d.add_heading("5. Acceptance checks", 1)
    d.add_paragraph("Criteria are engineering thresholds chosen for this prototype, not standards: "
                    f"LDR noise SD <= {CRIT['ldr_sd_max_pct_fs']}% of full scale; level separation >= "
                    f"{CRIT['ldr_sep_sd']:.0f} SD; DHT11 success >= {CRIT['dht_success_min']:.0f}%; PIR HIGH <= "
                    f"{CRIT['pir_still_max']:.0f}% when still and >= {CRIT['pir_motion_min']:.0f}% with motion; "
                    "logging period 90-110 ms.")
    table(d, ["Check", "Result", "Evidence"],
          [[c, "PASS" if ok is True else "FAIL" if ok is False else "NOT RUN", e] for c, ok, e in checks],
          [2.2, 0.8, 3.5])
    caption(d, "Table 5. Step S1.5 acceptance checks (computed automatically).")
    d.add_heading("6. Draft text for the thesis", 1)
    npass = sum(1 for _, ok, _ in checks if ok is True); nrun = sum(1 for _, ok, _ in checks if ok is not None)
    d.add_paragraph(f"All peripherals of the ESP32-S3 prototype were exercised at their assigned pins and logged at 10 Hz "
                    f"for {res['duration_s']:.0f} s. {npass} of {nrun} automated acceptance checks passed. The DHT11 "
                    f"returned valid readings in {h['success_pct']:.1f}% of {h['attempts']} attempts. "
                    + (f"The LDR showed {res['ldr_polarity']}, with a minimum separation of "
                       f"{res['ldr_separation_sd']:.0f} standard deviations between lighting levels. "
                       if "ldr_polarity" in res else "")
                    + "[Add one sentence on any FAIL above and the corrective action taken.]")
    d.add_heading("7. Limitations", 1)
    d.add_paragraph("Single session, single prototype, single room; no calibrated lux meter, so lighting levels are "
                    "qualitative labels rather than lux values. PIR behaviour depends on the sensitivity and delay "
                    "potentiometers, which should be recorded with the photo of the rig.")
    d.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("csv"); ap.add_argument("--out", default="reports/s1_5")
    a = ap.parse_args()
    os.makedirs(a.out, exist_ok=True)
    df, meta, kv = load(a.csv)
    res, checks, figs, labels = analyse(df, kv, a.out)
    docx_path = os.path.join(a.out, "S1_5_Report.docx")
    build_doc(docx_path, a.csv, df, meta, kv, res, checks, figs, labels)
    res["checks"] = [{"check": c, "pass": ok, "evidence": e} for c, ok, e in checks]
    res["source"] = kv.get("source", "unknown")
    json.dump(res, open(os.path.join(a.out, "s1_5_summary.json"), "w"), indent=2)
    for c, ok, e in checks:
        print(("PASS " if ok is True else "FAIL " if ok is False else "SKIP ") + c + " - " + e)
    print("Report:", docx_path)


if __name__ == "__main__":
    main()
